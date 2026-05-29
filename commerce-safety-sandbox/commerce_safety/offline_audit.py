from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .io import load_yaml, write_json, write_text


DATASET_FIELDS: dict[str, list[str]] = {
    "orders": [
        "order_id",
        "buyer_id",
        "sku",
        "quantity",
        "payment_status",
        "order_status",
        "created_at",
    ],
    "inventory": ["sku", "on_hand", "reserved", "warehouse"],
    "fulfillments": [
        "order_id",
        "sku",
        "quantity",
        "fulfillment_status",
        "tracking_number",
        "warehouse_status",
    ],
    "refunds": ["order_id", "amount", "refund_status", "reason", "approved_by"],
}

CRITICAL_WAREHOUSE_STATUSES = {
    "picked",
    "packed",
    "label_created",
    "carrier_scanned",
    "shipped",
}

POST_SHIPMENT_WAREHOUSE_STATUSES = {"carrier_scanned", "shipped", "delivered"}

SHIPPED_FULFILLMENT_STATUSES = {
    "fulfilled",
    "shipped",
    "carrier_scanned",
    "delivered",
}

REFUNDED_STATUSES = {"refunded", "issued", "completed", "succeeded", "paid"}
CANCELLED_ORDER_STATUSES = {"cancelled", "canceled", "cancel_requested"}
OPEN_ORDER_STATUSES = {"open", "paid", "processing", "unfulfilled"}
PAID_STATUSES = {"paid", "captured", "authorized"}


def make_audit_id() -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    return f"audit_{stamp}"


def normalize(value: Any) -> str:
    return str(value or "").strip()


def normalize_status(value: Any) -> str:
    return normalize(value).lower().replace(" ", "_").replace("-", "_")


def parse_int(value: Any, *, default: int = 0) -> int:
    text = normalize(value)
    if not text:
        return default
    try:
        return int(float(text))
    except ValueError:
        return default


def parse_float(value: Any, *, default: float = 0.0) -> float:
    text = normalize(value)
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return default


def redact_identifier(value: str) -> str:
    text = normalize(value)
    if not text:
        return ""
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()[:10]
    return f"buyer_{digest}"


def load_mapping(path: Path | None) -> dict[str, dict[str, str]]:
    if path is None:
        return {}
    data = load_yaml(path)
    mapping: dict[str, dict[str, str]] = {}
    for dataset, fields in data.items():
        if not isinstance(fields, dict):
            raise ValueError(f"Mapping for {dataset} must be a field mapping.")
        mapping[dataset] = {str(key): str(value) for key, value in fields.items()}
    return mapping


def read_csv_rows(path: Path) -> tuple[list[dict[str, str]], list[str], str | None]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        headers = list(reader.fieldnames or [])
        return [dict(row) for row in reader], headers, None


def read_xlsx_rows(
    path: Path, sheet_name: str | None
) -> tuple[list[dict[str, str]], list[str], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ValueError(
            "Reading .xlsx files requires openpyxl. Install it or provide CSV files."
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"Sheet {sheet_name!r} not found in {path}.")
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return [], [], worksheet.title
    headers = [normalize(value) for value in rows[0]]
    data_rows: list[dict[str, str]] = []
    for values in rows[1:]:
        if not any(normalize(value) for value in values):
            continue
        row = {
            header: normalize(values[index] if index < len(values) else "")
            for index, header in enumerate(headers)
            if header
        }
        data_rows.append(row)
    workbook.close()
    return data_rows, headers, worksheet.title


def read_table_rows(
    path: Path, sheet_name: str | None
) -> tuple[list[dict[str, str]], list[str], str | None]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix == ".xlsx":
        return read_xlsx_rows(path, sheet_name)
    raise ValueError(f"Unsupported input format for {path}. Use .csv or .xlsx.")


def canonicalize_rows(
    *,
    dataset: str,
    path: Path,
    mapping: dict[str, dict[str, str]],
    sheet_name: str | None = None,
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    raw_rows, headers, resolved_sheet_name = read_table_rows(path, sheet_name)
    fields = DATASET_FIELDS[dataset]
    dataset_mapping = mapping.get(dataset, {})
    missing_headers = [
        dataset_mapping.get(field, field)
        for field in fields
        if dataset_mapping.get(field, field) not in headers
    ]
    canonical_rows: list[dict[str, str]] = []
    missing_values: dict[str, int] = {field: 0 for field in fields}
    for raw_row in raw_rows:
        row: dict[str, str] = {}
        for field in fields:
            source_field = dataset_mapping.get(field, field)
            value = normalize(raw_row.get(source_field, ""))
            if field == "buyer_id":
                value = redact_identifier(value)
            row[field] = value
            if not value:
                missing_values[field] += 1
        canonical_rows.append(row)

    profile = {
        "source_path": str(path),
        "source_format": path.suffix.lower().lstrip("."),
        "sheet_name": resolved_sheet_name,
        "row_count": len(canonical_rows),
        "headers": headers,
        "missing_headers": missing_headers,
        "missing_values": {
            field: count for field, count in missing_values.items() if count
        },
    }
    return canonical_rows, profile


def write_redacted_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def build_data_quality(
    *,
    profiles: dict[str, dict[str, Any]],
    orders: list[dict[str, str]],
    inventory: list[dict[str, str]],
    fulfillments: list[dict[str, str]],
    refunds: list[dict[str, str]],
) -> dict[str, Any]:
    order_ids = {row["order_id"] for row in orders if row["order_id"]}
    inventory_skus = {row["sku"] for row in inventory if row["sku"]}
    order_skus = {row["sku"] for row in orders if row["sku"]}
    order_line_counter = defaultdict(int)
    for row in orders:
        key = (row["order_id"], row["sku"])
        if all(key):
            order_line_counter[key] += 1

    duplicate_order_lines = [
        {"order_id": order_id, "sku": sku, "row_count": count}
        for (order_id, sku), count in sorted(order_line_counter.items())
        if count > 1
    ]
    orphan_fulfillments = [
        {"order_id": row["order_id"], "sku": row["sku"]}
        for row in fulfillments
        if row["order_id"] and row["order_id"] not in order_ids
    ]
    orphan_refunds = [
        {"order_id": row["order_id"], "amount": row["amount"]}
        for row in refunds
        if row["order_id"] and row["order_id"] not in order_ids
    ]
    unknown_inventory_skus = sorted(order_skus - inventory_skus)
    negative_inventory = [
        {
            "sku": row["sku"],
            "warehouse": row["warehouse"],
            "on_hand": parse_int(row["on_hand"]),
            "reserved": parse_int(row["reserved"]),
        }
        for row in inventory
        if parse_int(row["on_hand"]) < 0 or parse_int(row["reserved"]) < 0
    ]
    return {
        "datasets": profiles,
        "issues": {
            "duplicate_order_lines": duplicate_order_lines,
            "orphan_fulfillments": orphan_fulfillments,
            "orphan_refunds": orphan_refunds,
            "unknown_inventory_skus": unknown_inventory_skus,
            "negative_inventory": negative_inventory,
        },
    }


def reconstruct_state(
    *,
    orders: list[dict[str, str]],
    inventory: list[dict[str, str]],
    fulfillments: list[dict[str, str]],
    refunds: list[dict[str, str]],
) -> dict[str, Any]:
    ordered_qty: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    orders_by_id: dict[str, dict[str, Any]] = {}
    for row in orders:
        order_id = row["order_id"]
        sku = row["sku"]
        quantity = parse_int(row["quantity"])
        if not order_id or not sku:
            continue
        ordered_qty[order_id][sku] += quantity
        order = orders_by_id.setdefault(
            order_id,
            {
                "order_id": order_id,
                "buyer_id": row["buyer_id"],
                "payment_status": row["payment_status"],
                "order_status": row["order_status"],
                "created_at": row["created_at"],
                "line_items": [],
            },
        )
        order["line_items"].append({"sku": sku, "quantity": quantity})

    inventory_by_sku: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"sku": "", "on_hand": 0, "reserved": 0, "warehouses": []}
    )
    for row in inventory:
        sku = row["sku"]
        if not sku:
            continue
        item = inventory_by_sku[sku]
        item["sku"] = sku
        on_hand = parse_int(row["on_hand"])
        reserved = parse_int(row["reserved"])
        item["on_hand"] += on_hand
        item["reserved"] += reserved
        item["warehouses"].append(
            {
                "warehouse": row["warehouse"],
                "on_hand": on_hand,
                "reserved": reserved,
                "available": on_hand - reserved,
            }
        )
    for item in inventory_by_sku.values():
        item["available"] = item["on_hand"] - item["reserved"]

    fulfilled_qty: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    fulfillment_rows: list[dict[str, Any]] = []
    for row in fulfillments:
        order_id = row["order_id"]
        sku = row["sku"]
        quantity = parse_int(row["quantity"])
        status = normalize_status(row["fulfillment_status"])
        warehouse_status = normalize_status(row["warehouse_status"])
        if status in SHIPPED_FULFILLMENT_STATUSES or warehouse_status in CRITICAL_WAREHOUSE_STATUSES:
            fulfilled_qty[order_id][sku] += quantity
        fulfillment_rows.append(
            {
                **row,
                "quantity": quantity,
                "fulfillment_status_normalized": status,
                "warehouse_status_normalized": warehouse_status,
            }
        )

    refund_rows: list[dict[str, Any]] = []
    refunds_by_order: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in refunds:
        refund = {
            **row,
            "amount": parse_float(row["amount"]),
            "refund_status_normalized": normalize_status(row["refund_status"]),
        }
        refund_rows.append(refund)
        if row["order_id"]:
            refunds_by_order[row["order_id"]].append(refund)

    return {
        "orders": orders_by_id,
        "inventory": dict(sorted(inventory_by_sku.items())),
        "fulfillments": fulfillment_rows,
        "refunds": refund_rows,
        "aggregates": {
            "ordered_qty": {
                order_id: dict(sorted(skus.items()))
                for order_id, skus in sorted(ordered_qty.items())
            },
            "fulfilled_qty": {
                order_id: dict(sorted(skus.items()))
                for order_id, skus in sorted(fulfilled_qty.items())
            },
            "refunds_by_order": {
                order_id: rows for order_id, rows in sorted(refunds_by_order.items())
            },
        },
    }


def finding(
    *,
    policy_id: str,
    severity: str,
    evidence: dict[str, Any],
    business_impact: str,
    recommendation: str,
) -> dict[str, Any]:
    return {
        "policy_id": policy_id,
        "severity": severity,
        "status": "failed",
        "evidence": evidence,
        "business_impact": business_impact,
        "recommendation": recommendation,
    }


def evaluate_offline_policies(state: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    orders = state["orders"]
    inventory = state["inventory"]
    ordered_qty = state["aggregates"]["ordered_qty"]
    fulfilled_qty = state["aggregates"]["fulfilled_qty"]
    refunds_by_order = state["aggregates"]["refunds_by_order"]

    for order_id, sku_qty in fulfilled_qty.items():
        for sku, fulfilled in sku_qty.items():
            ordered = ordered_qty.get(order_id, {}).get(sku, 0)
            if fulfilled > ordered:
                findings.append(
                    finding(
                        policy_id="offline_duplicate_fulfillment",
                        severity="critical",
                        evidence={
                            "order_id": order_id,
                            "sku": sku,
                            "ordered_quantity": ordered,
                            "fulfilled_quantity": fulfilled,
                        },
                        business_impact=(
                            "Historical fulfillment data shows more shipped quantity "
                            "than ordered quantity, which can mean duplicate shipment, "
                            "duplicate warehouse release, or an unsafe retry path."
                        ),
                        recommendation=(
                            "Add fulfillment idempotency and block duplicate warehouse "
                            "actions for the same order and SKU before the next automation change."
                        ),
                    )
                )

    for sku, item in inventory.items():
        if item["available"] < 0:
            findings.append(
                finding(
                    policy_id="offline_negative_available_inventory",
                    severity="critical",
                    evidence={
                        "sku": sku,
                        "on_hand": item["on_hand"],
                        "reserved": item["reserved"],
                        "available": item["available"],
                    },
                    business_impact=(
                        "Reserved inventory is greater than on-hand inventory. This "
                        "is a direct oversell signal and can lead to cancellations, "
                        "support load, and customer complaints."
                    ),
                    recommendation=(
                        "Require fresh inventory sync and successful reservation before "
                        "promising fulfillment or accepting automation-created commitments."
                    ),
                )
            )

    for order_id, order in orders.items():
        order_status = normalize_status(order["order_status"])
        payment_status = normalize_status(order["payment_status"])
        if order_status in OPEN_ORDER_STATUSES and payment_status in PAID_STATUSES:
            for line_item in order["line_items"]:
                sku = line_item["sku"]
                needed = int(line_item["quantity"])
                outstanding = needed - fulfilled_qty.get(order_id, {}).get(sku, 0)
                if outstanding <= 0:
                    continue
                available = int(inventory.get(sku, {}).get("available", 0))
                if available < outstanding:
                    findings.append(
                        finding(
                            policy_id="offline_paid_order_inventory_shortage",
                            severity="high",
                            evidence={
                                "order_id": order_id,
                                "sku": sku,
                                "outstanding_quantity": outstanding,
                                "available_inventory": available,
                            },
                            business_impact=(
                                "A paid open order does not have enough available "
                                "inventory in the reconstructed state. Automation can "
                                "promise stock that is no longer available."
                            ),
                            recommendation=(
                                "Hold paid orders that cannot be reserved and route them "
                                "to manual review before sending fulfillment promises."
                            ),
                        )
                    )

        if order_status in CANCELLED_ORDER_STATUSES:
            conflicting = [
                fulfillment
                for fulfillment in state["fulfillments"]
                if fulfillment["order_id"] == order_id
                and fulfillment["warehouse_status_normalized"]
                in CRITICAL_WAREHOUSE_STATUSES
            ]
            if conflicting:
                findings.append(
                    finding(
                        policy_id="offline_cancel_after_pick_pack_conflict",
                        severity="critical",
                        evidence={
                            "order_id": order_id,
                            "warehouse_statuses": [
                                item["warehouse_status"] for item in conflicting
                            ],
                            "refund_count": len(refunds_by_order.get(order_id, [])),
                        },
                        business_impact=(
                            "The order is cancelled while warehouse work is already "
                            "past the safe automatic-cancel point. This can cause refund "
                            "plus shipment, failed recall, or inventory confusion."
                        ),
                        recommendation=(
                            "When cancellation happens after pick or pack, create a hold "
                            "or warehouse cancellation request instead of automatic refund "
                            "and inventory release."
                        ),
                    )
                )

    shipped_orders = {
        fulfillment["order_id"]
        for fulfillment in state["fulfillments"]
        if fulfillment["fulfillment_status_normalized"] in SHIPPED_FULFILLMENT_STATUSES
        or fulfillment["warehouse_status_normalized"] in POST_SHIPMENT_WAREHOUSE_STATUSES
    }
    for order_id, refunds in refunds_by_order.items():
        issued_refunds = [
            refund
            for refund in refunds
            if refund["refund_status_normalized"] in REFUNDED_STATUSES
        ]
        unapproved = [refund for refund in issued_refunds if not refund["approved_by"]]
        if order_id in shipped_orders and unapproved:
            findings.append(
                finding(
                    policy_id="offline_refund_after_shipment_without_approval",
                    severity="high",
                    evidence={
                        "order_id": order_id,
                        "refund_amount": sum(refund["amount"] for refund in unapproved),
                        "unapproved_refund_count": len(unapproved),
                    },
                    business_impact=(
                        "Refunds were issued after shipment without an approval marker. "
                        "That creates money-plus-goods loss risk."
                    ),
                    recommendation=(
                        "Require approval for post-shipment refunds and block autonomous "
                        "refund issuance once warehouse or carrier state indicates shipment."
                    ),
                )
            )

    return findings


def severity_score(findings: list[dict[str, Any]]) -> int:
    weights = {"critical": 30, "high": 20, "medium": 10, "low": 5, "info": 1}
    return min(100, sum(weights.get(finding["severity"], 0) for finding in findings))


def build_offline_report(
    *,
    audit_id: str,
    data_quality: dict[str, Any],
    state: dict[str, Any],
    findings: list[dict[str, Any]],
) -> str:
    score = severity_score(findings)
    lines = [
        "# Offline Fulfillment Automation Audit",
        "",
        f"- Audit ID: `{audit_id}`",
        f"- Risk score: `{score}/100`",
        f"- Policy findings: `{len(findings)}`",
        "",
        "## Executive Summary",
        "",
    ]
    if findings:
        lines.extend(
            [
                "This export shows fulfillment automation risks that should be reviewed before a promotion, ERP workflow change, or AI automation launch.",
                "The biggest concern is not whether an API call can technically run; it is whether reconstructed commerce state shows money, inventory, refund, or warehouse conflicts.",
            ]
        )
    else:
        lines.append(
            "No policy findings were detected in this offline export. The data still needs operational review before production changes."
        )

    lines.extend(
        [
            "",
            "## Data Loaded",
            "",
        ]
    )
    for dataset, profile in data_quality["datasets"].items():
        lines.append(f"- {dataset}: `{profile['row_count']}` rows")

    issues = data_quality["issues"]
    lines.extend(
        [
            "",
            "## Data Quality Signals",
            "",
            f"- Duplicate order lines: `{len(issues['duplicate_order_lines'])}`",
            f"- Orphan fulfillments: `{len(issues['orphan_fulfillments'])}`",
            f"- Orphan refunds: `{len(issues['orphan_refunds'])}`",
            f"- Unknown inventory SKUs: `{len(issues['unknown_inventory_skus'])}`",
            f"- Negative inventory rows: `{len(issues['negative_inventory'])}`",
            "",
            "## Top Findings",
            "",
        ]
    )
    if not findings:
        lines.append("No high-risk policy findings were generated.")
    else:
        for item in findings:
            lines.extend(
                [
                    f"### {item['policy_id']}",
                    "",
                    f"- Severity: `{item['severity']}`",
                    f"- Business impact: {item['business_impact']}",
                    f"- Recommendation: {item['recommendation']}",
                    f"- Evidence: `{item['evidence']}`",
                    "",
                ]
            )

    lines.extend(
        [
            "## Reconstructed State Summary",
            "",
            f"- Orders reconstructed: `{len(state['orders'])}`",
            f"- SKUs reconstructed: `{len(state['inventory'])}`",
            f"- Fulfillment records: `{len(state['fulfillments'])}`",
            f"- Refund records: `{len(state['refunds'])}`",
            "",
            "## Recommended Fixes",
            "",
        ]
    )
    if findings:
        recommendations = []
        seen = set()
        for item in findings:
            if item["recommendation"] not in seen:
                recommendations.append(item["recommendation"])
                seen.add(item["recommendation"])
        for recommendation in recommendations:
            lines.append(f"- {recommendation}")
    else:
        lines.append("- Keep this audit as a baseline and rerun after automation changes.")

    lines.extend(
        [
            "",
            "## Saved Artifacts",
            "",
            "- `manifest.json`",
            "- `data_quality.json`",
            "- `state_reconstruction.json`",
            "- `policy_report.json`",
            "- `report.md`",
            "- `redacted_inputs/*.csv`",
        ]
    )
    return "\n".join(lines)


def run_offline_audit(
    *,
    orders_path: Path,
    inventory_path: Path,
    fulfillments_path: Path,
    refunds_path: Path,
    output_dir: Path,
    mapping_path: Path | None = None,
    sheet_names: dict[str, str | None] | None = None,
) -> dict[str, Any]:
    audit_id = make_audit_id()
    audit_path = output_dir / audit_id
    mapping = load_mapping(mapping_path)
    paths = {
        "orders": orders_path,
        "inventory": inventory_path,
        "fulfillments": fulfillments_path,
        "refunds": refunds_path,
    }
    rows: dict[str, list[dict[str, str]]] = {}
    profiles: dict[str, dict[str, Any]] = {}
    sheet_names = sheet_names or {}
    for dataset, path in paths.items():
        rows[dataset], profiles[dataset] = canonicalize_rows(
            dataset=dataset,
            path=path,
            mapping=mapping,
            sheet_name=sheet_names.get(dataset),
        )
        write_redacted_csv(
            audit_path / "redacted_inputs" / f"{dataset}.csv",
            rows[dataset],
            DATASET_FIELDS[dataset],
        )

    data_quality = build_data_quality(
        profiles=profiles,
        orders=rows["orders"],
        inventory=rows["inventory"],
        fulfillments=rows["fulfillments"],
        refunds=rows["refunds"],
    )
    state = reconstruct_state(
        orders=rows["orders"],
        inventory=rows["inventory"],
        fulfillments=rows["fulfillments"],
        refunds=rows["refunds"],
    )
    findings = evaluate_offline_policies(state)
    policy_report = {
        "audit_id": audit_id,
        "status": "failed" if findings else "passed",
        "risk_score": severity_score(findings),
        "findings": findings,
    }
    manifest = {
        "audit_id": audit_id,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "input_paths": {key: str(value) for key, value in paths.items()},
        "input_sheets": {
            key: value for key, value in sheet_names.items() if value is not None
        },
        "mapping_path": str(mapping_path) if mapping_path else None,
        "artifacts": [
            "manifest.json",
            "data_quality.json",
            "state_reconstruction.json",
            "policy_report.json",
            "report.md",
            "redacted_inputs/orders.csv",
            "redacted_inputs/inventory.csv",
            "redacted_inputs/fulfillments.csv",
            "redacted_inputs/refunds.csv",
        ],
    }
    report = build_offline_report(
        audit_id=audit_id,
        data_quality=data_quality,
        state=state,
        findings=findings,
    )

    write_json(audit_path / "manifest.json", manifest)
    write_json(audit_path / "data_quality.json", data_quality)
    write_json(audit_path / "state_reconstruction.json", state)
    write_json(audit_path / "policy_report.json", policy_report)
    write_text(audit_path / "report.md", report)
    return {
        "audit_id": audit_id,
        "audit_path": str(audit_path),
        "status": policy_report["status"],
        "risk_score": policy_report["risk_score"],
        "findings": findings,
    }
